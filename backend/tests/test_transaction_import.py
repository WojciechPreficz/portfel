import unittest
from datetime import datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import Workbook

from app.services.transaction_import import read_bossa_purchases, read_deposits, read_purchases


class TransactionImportTest(unittest.TestCase):
    def test_bossa_csv_reads_transactions(self):
        content = (
            "Data;Walor;Rachunek;Waluta;Liczba;Strona;Cena;Wartość przed prowizją;Prowizja;Wartość po prowizji\n"
            "10.10.2022 09:04:13;Vanguard LifeStrategy 80% Equity UCITS ETF;IKE 816742;PLN;132,00;K;130,307;17200,49;49,88;17250,37\n"
            "11.10.2022 09:04:13;Vanguard LifeStrategy 80% Equity UCITS ETF;IKE 816742;PLN;2,00;S;140,000;280,00;0,00;280,00\n"
        ).encode()

        purchases, errors = read_bossa_purchases(content)

        self.assertEqual(errors, [])
        self.assertEqual(purchases[0]["date"].isoformat(), "2022-10-10")
        self.assertEqual(purchases[0]["quantity"], Decimal("132.00"))
        self.assertEqual(purchases[0]["price"], Decimal("130.307"))
        self.assertEqual(purchases[0]["commission"], Decimal("49.88"))
        self.assertEqual(purchases[0]["type"], "BUY")
        self.assertEqual(purchases[1]["type"], "SELL")

    def test_xstation_cash_operations_reads_deposits(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Cash Operations"
        sheet.append(["Time", "Type", "Ticker", "Comment", "Amount"])
        sheet.append([datetime(2026, 9, 11), "Deposit", None, "Cash deposit", 35000])
        sheet.append([datetime(2026, 9, 12), "Deposit", None, "Cash deposit", -5000])

        content = BytesIO()
        workbook.save(content)

        deposits, errors = read_deposits(content.getvalue())

        self.assertEqual(errors, [])
        self.assertEqual([deposit["amount"] for deposit in deposits], [Decimal("35000"), Decimal("5000")])
        self.assertEqual([deposit["currency"] for deposit in deposits], ["PLN", "PLN"])

    def test_xstation_cash_operations_uses_ticker_and_instrument_name(self):
        workbook = Workbook()
        workbook.active.title = "Summary"
        sheet = workbook.create_sheet("Cash Operations")
        sheet.append(["Time", "Type", "Ticker", "Instrument", "Comment", "Amount"])
        sheet.append([datetime(2026, 9, 11), "Buy", "ABC.US", "Example Corp", "OPEN BUY 2/95 @ 85.20", -170.40])
        sheet.append([datetime(2026, 9, 12), "Buy", "ABC.US", "Example Corp", "OPEN BUY 93/95 @ 85.40", -7942.20])

        content = BytesIO()
        workbook.save(content)

        purchases, errors = read_purchases(content.getvalue())

        self.assertEqual(errors, [])
        self.assertEqual([purchase["ticker"] for purchase in purchases], ["ABC", "ABC"])
        self.assertEqual([purchase["name"] for purchase in purchases], ["Example Corp", "Example Corp"])
        self.assertEqual([purchase["quantity"] for purchase in purchases], [Decimal("2"), Decimal("93")])
        self.assertEqual([purchase["price"] for purchase in purchases], [Decimal("85.20"), Decimal("85.40")])

    def test_xstation_uses_comment_quantity_instead_of_amount(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Time", "Type", "Symbol", "Comment", "Amount", "Price"])
        sheet.append([datetime(2026, 9, 11), "Buy", "WEC.US", "Buy 7/18 @ 106.05", 18, 106.05])

        content = BytesIO()
        workbook.save(content)

        purchases, errors = read_purchases(content.getvalue())

        self.assertEqual(errors, [])
        self.assertEqual(purchases[0]["quantity"], 7)
        self.assertEqual(purchases[0]["price"], Decimal("106.05"))

    def test_xstation_uses_comment_quantity_when_amount_is_negative(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Time", "Type", "Symbol", "Comment", "Amount", "Price"])
        sheet.append([datetime(2026, 9, 11), "Buy", "WEC.US", "Buy 7/18 @ 106.05", -18, 106.05])

        content = BytesIO()
        workbook.save(content)

        purchases, errors = read_purchases(content.getvalue())

        self.assertEqual(errors, [])
        self.assertEqual(purchases[0]["quantity"], 7)

    def test_broadcom_purchase_is_adjusted_for_ten_to_one_split(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Time", "Type", "Symbol", "Comment", "Amount", "Price"])
        sheet.append([datetime(2023, 10, 31), "Buy", "AVGO.US", "Buy 1 @ 837.92", 1, 837.92])

        content = BytesIO()
        workbook.save(content)

        purchases, errors = read_purchases(content.getvalue())

        self.assertEqual(errors, [])
        self.assertEqual(purchases[0]["quantity"], Decimal("10"))
        self.assertEqual(purchases[0]["price"], Decimal("83.792"))
        self.assertEqual(purchases[0]["date"], datetime(2023, 10, 31).date())


if __name__ == "__main__":
    unittest.main()