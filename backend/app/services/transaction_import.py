from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
import csv
import re

from openpyxl import load_workbook


HEADER_ALIASES = {
    "date": {"data", "data transakcji", "data operacji", "data zawarcia", "data realizacji", "time", "open time utc", "trade date", "transaction date"},
    "ticker": {"ticker", "symbol", "kod", "walor", "nazwa skrocona"},
    "instrument_name": {"instrument", "instruments", "nazwa", "nazwa instrumentu"},
    "isin": {"isin"},
    "quantity": {"ilosc", "liczba", "liczba sztuk", "ilosc instrumentow", "quantity", "qty", "wolumen"},
    "price": {"cena", "kurs", "kurs transakcji", "cena transakcji", "price", "cena jednostkowa", "kurs jednostkowy"},
    "value": {"wartosc", "wartosc transakcji", "wartosc brutto", "wartosc netto", "kwota", "amount", "value", "transaction value"},
    "currency": {"waluta", "currency"},
    "commission": {"prowizja", "commission", "oplaty", "fee"},
    "type": {"typ", "typ transakcji", "operacja", "rodzaj operacji", "type"},
    "comment": {"komentarz", "comment"},
    "category": {"kategoria", "category"},
}

STOCK_SPLITS = {
    "AVGO": ((date(2024, 7, 15), Decimal("10")),),
}


def _normalise(value: object) -> str:
    text = str(value or "").strip().lower()
    for source, target in (("ł", "l"), ("ą", "a"), ("ę", "e"), ("ó", "o"), ("ś", "s"), ("ź", "z"), ("ż", "z"), ("ć", "c"), ("ń", "n")):
        text = text.replace(source, target)
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def _decimal(value: object, field: str, row_number: int) -> Decimal:
    if value is None or str(value).strip() == "":
        raise ValueError(f"wiersz {row_number}: brak pola {field}")
    text = str(value).strip().replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"wiersz {row_number}: nieprawidlowa wartosc {field}: {value}") from exc


def _date(value: object, row_number: int) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for pattern in ("%Y-%m-%d", "%d.%m.%Y %H:%M:%S", "%d.%m.%Y", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(str(value).strip(), pattern).date()
        except ValueError:
            continue
    raise ValueError(f"wiersz {row_number}: nieprawidlowa data: {value}")


def _normalise_ticker(value: object) -> str | None:
    if not value:
        return None
    ticker = str(value).strip().upper()
    for suffix in (".US", ".PL"):
        if ticker.endswith(suffix):
            return ticker[: -len(suffix)]
    return ticker


def _apply_stock_splits(purchase: dict) -> dict:
    ticker = purchase["ticker"]
    for effective_date, ratio in STOCK_SPLITS.get(ticker, ()):
        if purchase["date"] < effective_date:
            purchase["quantity"] *= ratio
            purchase["price"] /= ratio
    return purchase


def read_purchases(content: bytes) -> tuple[list[dict], list[str]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    header_candidates = []
    xstation_sheet = False
    for sheet_index, sheet in enumerate(workbook.worksheets):
        if _normalise(sheet.title) == "cash operations":
            xstation_sheet = True
        sheet.reset_dimensions()
        rows = list(sheet.iter_rows(values_only=True))
        if xstation_sheet:
            header_candidates = []
        for row_index, row in enumerate(rows):
            normalised_values = {_normalise(value) for value in row if value not in (None, "")}
            recognised = sum(
                any(_normalise(value) in aliases for aliases in HEADER_ALIASES.values())
                for value in row
                if value not in (None, "")
            )
            if {"comment", "amount", "time", "type"}.issubset(normalised_values):
                recognised += 5
            if recognised:
                header_candidates.append((recognised, -sheet_index, -row_index, rows))
        if xstation_sheet:
            break
    if not header_candidates:
        raise ValueError("Arkusz jest pusty")
    _, _, negative_header_index, rows = max(header_candidates)
    header_index = -negative_header_index

    headers = {}
    for index, value in enumerate(rows[header_index]):
        normalised = _normalise(value)
        for field, aliases in HEADER_ALIASES.items():
            if normalised in aliases:
                headers[field] = index
                break
    missing = [field for field in ("date", "quantity") if field not in headers]
    if "price" not in headers and "value" not in headers:
        missing.append("price lub value")
    if "ticker" not in headers and "isin" not in headers:
        missing.append("ticker lub isin")
    xstation_cash = {"date", "ticker", "type", "comment"}.issubset(headers)
    if xstation_cash:
        missing = [field for field in missing if field not in {"quantity", "price lub value"}]
    if missing:
        raise ValueError(f"Brak wymaganych kolumn: {', '.join(missing)}")

    purchases = []
    errors = []
    for row_number, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        if not any(value not in (None, "") for value in row):
            continue
        operation = _normalise(row[headers["type"]]) if "type" in headers else "buy"
        if operation and not any(word in operation for word in ("buy", "kup", "zakup", "purchase", "naby")):
            continue
        try:
            if xstation_cash:
                match = re.search(r"\bbuy\s+(\d+(?:/\d+)?)\s*@\s*([\d.,]+)", str(row[headers["comment"]]), re.IGNORECASE)
                if not match:
                    raise ValueError(f"wiersz {row_number}: nie rozpoznano ilosci i ceny w komentarzu")
                quantity_text = match.group(1)
                if "/" in quantity_text:
                    quantity = Decimal(quantity_text.split("/", 1)[0])
                else:
                    quantity = Decimal(quantity_text)
                price = _decimal(match.group(2), "price", row_number)
            else:
                quantity = _decimal(row[headers["quantity"]], "quantity", row_number)
                if quantity <= 0:
                    raise ValueError(f"wiersz {row_number}: quantity musi byc > 0")
                price = _decimal(row[headers["price"]], "price", row_number) if "price" in headers else None
                if price is None:
                    price = _decimal(row[headers["value"]], "value", row_number) / quantity
            if quantity <= 0:
                raise ValueError(f"wiersz {row_number}: quantity musi byc > 0")
            if price < 0:
                raise ValueError(f"wiersz {row_number}: price musi byc >= 0")
            purchases.append(_apply_stock_splits({
                "row_number": row_number,
                "date": _date(row[headers["date"]], row_number),
                "ticker": _normalise_ticker(row[headers["ticker"]]) if "ticker" in headers else None,
                "raw_ticker": str(row[headers["ticker"]]).strip().upper() if "ticker" in headers and row[headers["ticker"]] else None,
                "name": str(row[headers["instrument_name"]]).strip() if "instrument_name" in headers and row[headers["instrument_name"]] else None,
                "category": _normalise(row[headers["category"]]) if "category" in headers and row[headers["category"]] else None,
                "isin": str(row[headers["isin"]]).strip().upper() if "isin" in headers and row[headers["isin"]] else None,
                "quantity": quantity,
                "price": price,
                "currency": str(row[headers["currency"]]).strip().upper() if "currency" in headers and row[headers["currency"]] else None,
                "commission": _decimal(row[headers["commission"]], "commission", row_number) if "commission" in headers and row[headers["commission"]] not in (None, "") else Decimal("0"),
            }))
        except ValueError as exc:
            errors.append(str(exc))
    return purchases, errors


def read_bossa_purchases(content: bytes) -> tuple[list[dict], list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("cp1250")
    reader = csv.DictReader(StringIO(text), delimiter=";")
    headers = {_normalise(header): header for header in (reader.fieldnames or [])}
    required = {"data", "walor", "waluta", "liczba", "strona", "cena", "prowizja"}
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(f"Brak wymaganych kolumn: {', '.join(missing)}")

    purchases = []
    errors = []
    for row_number, row in enumerate(reader, start=2):
        if not any(value and value.strip() for value in row.values()):
            continue
        try:
            side = _normalise(row[headers["strona"]])
            if side not in {"k", "s"}:
                raise ValueError(f"wiersz {row_number}: nieprawidlowa strona transakcji: {row[headers['strona']]}")
            quantity = _decimal(row[headers["liczba"]], "quantity", row_number)
            price = _decimal(row[headers["cena"]], "price", row_number)
            commission = _decimal(row[headers["prowizja"]], "commission", row_number)
            if quantity <= 0:
                raise ValueError(f"wiersz {row_number}: quantity musi byc > 0")
            if price < 0:
                raise ValueError(f"wiersz {row_number}: price musi byc >= 0")
            if commission < 0:
                raise ValueError(f"wiersz {row_number}: commission musi byc >= 0")
            name = row[headers["walor"]].strip()
            if not name:
                raise ValueError(f"wiersz {row_number}: brak pola instrument")
            purchases.append({
                "row_number": row_number,
                "date": _date(row[headers["data"]], row_number),
                "ticker": None,
                "raw_ticker": None,
                "name": name,
                "category": "etf",
                "isin": None,
                "quantity": quantity,
                "price": price,
                "currency": row[headers["waluta"]].strip().upper(),
                "commission": commission,
                "type": "BUY" if side == "k" else "SELL",
            })
        except ValueError as exc:
            errors.append(str(exc))
    return purchases, errors


def read_deposits(content: bytes) -> tuple[list[dict], list[str]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = next((sheet for sheet in workbook.worksheets if _normalise(sheet.title) == "cash operations"), None)
    if sheet is None:
        return [], []

    sheet.reset_dimensions()
    rows = list(sheet.iter_rows(values_only=True))
    header_index = None
    headers = {}
    for row_index, row in enumerate(rows):
        candidate = {}
        for index, value in enumerate(row):
            normalised = _normalise(value)
            for field, aliases in HEADER_ALIASES.items():
                if normalised in aliases:
                    candidate[field] = index
                    break
        if "date" in candidate and "type" in candidate and "value" in candidate:
            header_index = row_index
            headers = candidate
            break
    if header_index is None:
        return [], []

    deposits = []
    errors = []
    for row_number, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        operation = _normalise(row[headers["type"]]) if len(row) > headers["type"] else ""
        if operation not in {"deposit", "wplata", "cash deposit"}:
            continue
        try:
            amount = abs(_decimal(row[headers["value"]], "amount", row_number))
            if amount <= 0:
                raise ValueError(f"wiersz {row_number}: amount musi byc > 0")
            deposits.append(
                {
                    "date": _date(row[headers["date"]], row_number),
                    "amount": amount,
                    "currency": (
                        str(row[headers["currency"]]).strip().upper()
                        if "currency" in headers and row[headers["currency"]]
                        else "PLN"
                    ),
                }
            )
        except ValueError as exc:
            errors.append(str(exc))
    return deposits, errors